from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import call, patch

from openpyxl import Workbook

from app.services import financial_pusher


class FakeQuery:
    def __init__(self, table: str, calls: list[dict]):
        self.table = table
        self.calls = calls
        self.operation = ""
        self.payload = None
        self.filters = []

    def select(self, columns):
        self.operation = "select"
        self.payload = columns
        return self

    def update(self, payload):
        self.operation = "update"
        self.payload = payload
        return self

    def insert(self, payload):
        self.operation = "insert"
        self.payload = payload
        return self

    def delete(self):
        self.operation = "delete"
        return self

    def eq(self, column, value):
        self.filters.append(("eq", column, value))
        return self

    def is_(self, column, value):
        self.filters.append(("is", column, value))
        return self

    def in_(self, column, values):
        self.filters.append(("in", column, values))
        return self

    def execute(self):
        call = {
            "table": self.table,
            "operation": self.operation,
            "payload": self.payload,
            "filters": self.filters,
        }
        self.calls.append(call)
        data = [{"id": "existing-b306-line"}] if self.table == "iafr_line_items" and self.operation == "select" else [{}]
        return SimpleNamespace(data=data)


class FinancialPusherPatchTest(TestCase):
    def test_patch_updates_only_selected_account_across_all_months(self):
        calls = []
        accounts = [
            {
                "id": "b306-account",
                "account_code": "B.3.06",
                "account_name": "Charge Over / Above",
                "section_code": "B",
                "subsection_code": "other_receipts",
                "account_type": "receipt",
            }
        ]
        rows = [
            {
                "id": "push-row",
                "institution_id": "parish-id",
                "parish_code": "P-1",
                "parish_name": "Test Parish",
                "reporting_year": 2025,
                "reporting_month": 12,
                "validation_status": "ready",
                "source_row_number": 7,
                "raw_values": {"charge-key": 123.45},
                "cleaned_values": {"charge-key": 123.45},
                "issues": [],
            },
            {
                "id": "push-row-november",
                "institution_id": "parish-id",
                "parish_code": "P-1",
                "parish_name": "Test Parish",
                "reporting_year": 2025,
                "reporting_month": 11,
                "validation_status": "ready",
                "source_row_number": 7,
                "raw_values": {"charge-key": 999.99},
                "cleaned_values": {"charge-key": 999.99},
                "issues": [],
            },
        ]
        mappings = [
            {
                "key": "charge-key",
                "action": "map",
                "canonicalAccountCode": "B.3.06",
                "sourceHeader": "Charge Over/Above",
            }
        ]

        def fake_table(_schema, table):
            return FakeQuery(table, calls)

        with (
            patch.object(financial_pusher, "list_accounts", return_value=accounts),
            patch.object(financial_pusher, "_load_push_rows_for_commit", return_value=rows),
            patch.object(financial_pusher, "_get_existing_record", return_value="financial-record") as existing_record,
            patch.object(financial_pusher, "get_table", side_effect=fake_table),
        ):
            result = financial_pusher.commit_batch(
                "batch",
                mappings,
                import_mode="patch_selected",
                target_year=2025,
                target_month=None,
            )

        line_calls = [call for call in calls if call["table"] == "iafr_line_items"]
        self.assertFalse(any(call["operation"] == "delete" for call in line_calls))
        updates = [call for call in line_calls if call["operation"] == "update"]
        self.assertEqual(len(updates), 2)
        self.assertEqual([update["payload"]["amount"] for update in updates], [123.45, 999.99])
        self.assertTrue(all(("eq", "id", "existing-b306-line") in update["filters"] for update in updates))
        self.assertEqual(
            existing_record.call_args_list,
            [
                call("parish-id", 2025, "Dec"),
                call("parish-id", 2025, "Nov"),
            ],
        )
        self.assertEqual(result["committedRows"], 2)
        self.assertEqual(result["lineItemsCreated"], 2)
        self.assertEqual(result["targetYear"], 2025)
        self.assertIsNone(result["targetMonth"])

    def test_2025_unlabeled_sacrament_amount_is_recovered(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "December 2025"
        sheet.merge_cells("C3:J3")
        sheet["C3"] = "Sacraments"
        sheet.merge_cells("D4:J4")
        sheet["D4"] = "Baptism (Infant)"
        sheet["D6"] = "Rate Amount"
        sheet["E5"] = "Quantity"
        sheet["E6"] = "GRATIS"
        sheet["F6"] = "CHARGEABLE"
        # G5/G6 were accidentally left blank in the real 2025 template.
        sheet["H5"] = "Over/Above Amount"
        sheet["I5"] = "Total Amount as Over/Above"
        sheet["J5"] = "Total Amount"
        sheet["A7"] = "D1-1"
        sheet["B7"] = "Test Parish"
        sheet["G7"] = 1600

        columns = financial_pusher._source_columns(sheet, {}, 2025)
        repaired = next(column for column in columns if column.column == "G")

        self.assertEqual(repaired.source_header, "Total Amount as Prescribed")
        self.assertEqual(repaired.suggested_account_code, "A.1.02.01")
        self.assertEqual(repaired.suggested_field, "sacrament_prescribed_total")
