"""Historical parish financial PUSHER parser and commit helpers.

This module handles consolidated annual workbooks where each monthly sheet
contains many parish rows. It stages a reviewable interpretation first, then
commits approved mappings into the canonical parish finance tables.
"""

from __future__ import annotations

import base64
import hashlib
import io
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from app.services.supabase_client import get_table

MONTHS = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]
MONTH_ABBREV = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH_TO_NUM = {month.lower(): i + 1 for i, month in enumerate(MONTHS)}
MONTH_TO_SHORT = {i + 1: short for i, short in enumerate(MONTH_ABBREV)}
FORMULA_ERRORS = {"#N/A", "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!"}
MONTH_NAME_PATTERN = "|".join(
    sorted(
        {month.lower() for month in MONTHS} | {month.lower() for month in MONTH_ABBREV},
        key=len,
        reverse=True,
    )
)
MONTH_DAY_RE = re.compile(
    rf"\b(?:{MONTH_NAME_PATTERN})\.?\s+\d{{1,2}}(?:st|nd|rd|th)?\b"
    rf"|\b\d{{1,2}}(?:st|nd|rd|th)?\s+(?:{MONTH_NAME_PATTERN})\.?\b",
    re.IGNORECASE,
)


@dataclass
class SourceColumn:
    key: str
    column: str
    index: int
    source_section: str
    source_header: str
    header_path: str
    suggested_account_code: str | None
    suggested_account_name: str | None
    suggested_field: str | None
    confidence: float
    status: str
    reason: str
    aggregation_rule: str = "sum"
    is_combined: bool = False


def _source_column_key(parts: list[str], occurrence: int) -> str:
    text = " > ".join(part.strip() for part in parts if part and part.strip())
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text.lower())
        if not unicodedata.combining(char)
    )
    text = re.sub(r"[^a-z0-9]+", " ", text)
    base = re.sub(r"\s+", " ", text).strip() or "column"
    return base if occurrence == 1 else f"{base} #{occurrence}"


def normalize_name(value: str | None) -> str:
    text = (value or "").lower()
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    text = re.sub(r"\bsto\b\.?", "santo", text)
    text = re.sub(r"\bsta\b\.?", "santa", text)
    text = re.sub(r"\bsts\b\.?", "saint", text)
    text = re.sub(r"\bst\b\.?", "saint", text)
    text = re.sub(r"\bsaints\b", "saint", text)
    replacements = {
        "san": "saint",
        "santo": "saint",
        "santa": "saint",
        "juan": "john",
        "bautista": "baptist",
        "ebanghelista": "evangelist",
        "evangelista": "evangelist",
        "santiago": "james",
        "pedro": "peter",
        "pablo": "paul",
        "maria": "mary",
        "magdalena": "magdalene",
        "apostol": "apostle",
        "arkanghel": "archangel",
        "nuestra": "our",
        "senora": "lady",
    }
    for source, target in replacements.items():
        text = re.sub(rf"\b{source}\b", target, text)
    text = text.replace("rosaryo", "rosario")
    text = text.replace("sepulchre", "sepulcher")
    text = re.sub(r"\b(parish|parokya|parroquia|diocesan shrine of|shrine of|the|de|del|dela|delos|ng)\b", " ", text)
    text = re.sub(r"\bat\b", "and", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _name_similarity(source: str, candidate: str) -> float:
    base = SequenceMatcher(None, source, candidate).ratio()
    source_tokens = set(source.split())
    candidate_tokens = set(candidate.split())
    common = {"saint", "our", "lady", "of", "and", "y"}
    source_sig = {token for token in source_tokens if token not in common}
    candidate_sig = {token for token in candidate_tokens if token not in common}
    if source_sig and candidate_sig:
        shorter, longer = sorted([source_sig, candidate_sig], key=len)
        if len(shorter) >= 2 and shorter.issubset(longer):
            base = max(base, 0.94)
        overlap = len(source_sig & candidate_sig)
        if overlap:
            base = max(base, (2 * overlap) / (len(source_sig) + len(candidate_sig)))
    return base


def parse_money(value: Any) -> tuple[float, bool]:
    if value is None or value == "":
        return 0.0, False
    if isinstance(value, (int, float)):
        return float(value), False
    text = str(value).strip()
    if text.upper() in FORMULA_ERRORS:
        return 0.0, True
    negative = text.startswith("(") and text.endswith(")")
    cleaned = text.replace("PHP", "").replace("Php", "").replace("₱", "").replace(",", "").strip()
    cleaned = cleaned.strip("()")
    try:
        number = float(cleaned)
    except ValueError:
        return 0.0, True
    return (-number if negative else number), False


def _month_year_from_sheet(sheet_name: str) -> tuple[int, int] | None:
    match = re.match(r"^([A-Za-z]+)\s+(\d{4})$", sheet_name.strip())
    if not match:
        return None
    month = MONTH_TO_NUM.get(match.group(1).lower())
    return (month, int(match.group(2))) if month else None


def _header_grid(ws, max_header_row: int = 6) -> list[list[Any]]:
    grid = [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)] for r in range(1, max_header_row + 1)]
    for merged in ws.merged_cells.ranges:
      if merged.min_row > max_header_row:
          continue
      value = ws.cell(merged.min_row, merged.min_col).value
      for r in range(merged.min_row, min(merged.max_row, max_header_row) + 1):
          for c in range(merged.min_col, merged.max_col + 1):
              grid[r - 1][c - 1] = value
    return grid


def _clean_header(value: Any) -> str:
    text = str(value or "").strip()
    text = text.replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text


def _has_month_day(value: str) -> bool:
    return bool(MONTH_DAY_RE.search(value))


def _is_financial_value_column(parts: list[str]) -> bool:
    joined = " ".join(parts).lower()
    if not joined:
        return False
    if parts and parts[0].lower() == "sacraments" and parts[-1].lower() == "rate":
        return True
    if any(term in joined for term in ["rate amount", "gratis", "chargeable"]):
        return True
    if "quantity" in joined:
        return any(term in joined for term in ["gratis", "chargeable"])
    if joined.endswith(" rate") and "tax rate" not in joined and "prescribed rate" not in joined:
        return False
    include_terms = [
        "amount", "collection", "collections", "donation", "donations", "income", "subsidy",
        "receipts", "expenses", "borrowings", "balance", "stipend", "share", "bill", "wages",
        "contributions", "compensations", "allowances", "labor", "materials", "net receipts",
    ]
    return "tax rate" in joined or any(term in joined for term in include_terms)


def _sacrament_account(text: str) -> tuple[str | None, str | None]:
    if (
        "confirmation" in text
        and "minister" not in text
        and "bishop" not in text
        and "pension" not in text
        and "progressive" not in text
    ):
        return "A.2.01", "Confirmation"
    if "baptism (infant)" in text or ("baptism" in text and "infant" in text):
        return "A.1.02", "Baptism - Infant"
    if "baptism (adult)" in text or ("baptism" in text and "adult" in text):
        return "A.1.03", "Baptism - Adult"
    if "baptism" in text:
        return "A.1.01", "Baptism - General / Regular"
    if "wedding with mass" in text:
        return "A.1.04", "Wedding with Mass"
    if "wedding w/out mass" in text or "wedding without mass" in text:
        return "A.1.05", "Wedding without Mass"
    if "funeral mass" in text:
        return "A.1.06", "Funeral Mass"
    if "funeral blessings" in text:
        return "A.1.07", "Funeral Blessings"
    if "certificates" in text:
        return "A.1.08", "Certificates"
    if "marriage banns" in text:
        return "A.1.09", "Marriage Banns"
    if "permits" in text:
        return "A.1.10", "Permits"
    return None, None


def _sacrament_breakdown_account(base_code: str | None, kind: str) -> str | None:
    if not base_code or not (base_code.startswith("A.1.") or base_code == "A.2.01"):
        return None
    suffix = "01" if kind == "prescribed" else "02"
    return f"{base_code}.{suffix}"


def _suggest_mapping(parts: list[str], account_names: dict[str, str]) -> tuple[str | None, str | None, str | None, float, str, str, str, bool]:
    text = " ".join(parts).lower()
    section = parts[0].lower() if parts else ""
    account: str | None = None
    field: str | None = None
    confidence = 0.0
    status = "suggested"
    reason = "Needs finance review before import."
    rule = "sum"
    combined = False

    sacrament_code, sacrament_name = _sacrament_account(text)

    if section == "sacraments" and parts and parts[-1].lower() == "rate":
        return None, None, "sacrament_diocese_share_rate", 0.9, "needs_review", "Sacrament diocese-share rate; saved as memo/checking detail, not posted as money.", "memo", False
    if sacrament_code and ("rate amount" in text or "prescribed rate" in text):
        account, field, confidence, status, rule = sacrament_code, "sacrament_prescribed_rate", 0.9, "needs_review", "memo"
        reason = f"Prescribed rate for {sacrament_name}; saved as memo/checking detail, not posted as money."
    elif sacrament_code and "gratis" in text and ("quantity" in text or "chargeable" not in text):
        account, field, confidence, status, rule = sacrament_code, "sacrament_gratis_quantity", 0.9, "needs_review", "memo"
        reason = f"Gratis quantity for {sacrament_name}; saved as count detail, not posted as money."
    elif sacrament_code and "chargeable" in text and "quantity" in text:
        account, field, confidence, status, rule = sacrament_code, "sacrament_chargeable_quantity", 0.9, "needs_review", "memo"
        reason = f"Chargeable quantity for {sacrament_name}; saved as count detail, not posted as money."
    elif "quantity" in text:
        return None, None, None, 1.0, "ignored", "Non-amount helper column.", "ignore", False
    elif "tax rate" in text:
        account, field, confidence, status, rule = "B.1.04", "tax_rate", 0.98, "approved", "memo"
        reason = "Mass collection tax rate; saved as a checking/memo value."
    elif "remittances" in section and _has_month_day(text):
        account, confidence = "F.3.01", 0.95
        reason = "Dated remittance/special collection column maps to Special Collections."
    elif sacrament_code and "total amount" in text and "prescribed" not in text and "over/above" not in text:
        account, field, confidence, status, rule = sacrament_code, "sacrament_total_check", 0.9, "needs_review", "memo"
        reason = f"Final sacrament total for {sacrament_name}; kept as checking detail because prescribed and over/above totals are mapped separately."
    elif sacrament_code and "total amount as over/above" in text:
        account = _sacrament_breakdown_account(sacrament_code, "over_above") or "A.1.11"
        field, confidence, status = "sacrament_over_above_total", 0.92, "needs_review"
        reason = f"Total over/above amount maps to the {sacrament_name} over/above breakdown account."
    elif sacrament_code and ("charge over" in text or "over/above" in text):
        account, field, confidence, status, rule = "A.1.11", "sacrament_over_above_amount", 0.86, "needs_review", "memo"
        reason = "Over/above charge detail; saved as checking detail, not posted as a separate receipt."
    elif sacrament_code and "total amount as prescribed" in text:
        account = _sacrament_breakdown_account(sacrament_code, "prescribed") or sacrament_code
        field, confidence, status = "sacrament_prescribed_total", 0.92, "needs_review"
        reason = f"Total prescribed amount maps to the {sacrament_name} prescribed breakdown account."
    elif "internet" in text and ("telephone" in text or "communication" in text):
        account, confidence, combined = "D.4.05", 0.98, True
        reason = "Combined telephone/internet/communications expense column."
    elif "telephone" in text:
        account, confidence = "D.4.01", 0.98
        reason = "Telephone expense column."
    elif "cable" in text:
        account, confidence = "D.4.02", 0.98
        reason = "Cable expense column."
    elif "internet" in text:
        account, confidence = "D.4.03", 0.98
        reason = "Internet expense column."
    elif "mailing" in text:
        account, confidence = "D.4.04", 0.98
        reason = "Mailing expense column."
    elif "electric" in text or "water bill" in text:
        account, confidence = ("D.3.01", 0.98) if "electric" in text else ("D.3.02", 0.98)
        reason = "Utility expense column."
    elif "sss" in text:
        account, confidence = "D.2.01", 0.98
        reason = "Government contribution column."
    elif "hdmf" in text or "pag-ibig" in text:
        account, confidence = "D.2.02", 0.98
        reason = "Government contribution column."
    elif "philhealth" in text or "phic" in text:
        account, confidence = "D.2.03", 0.98
        reason = "Government contribution column."
    elif any(k in text for k in ["salaries", "wages", "remuneration", "compensations", "allowances", "13th month", "bonuses"]):
        if "salaries" in text or "wages" in text:
            account = "D.1.01"
        elif "remuneration" in text:
            account = "D.1.02"
        elif "13th month" in text or "bonuses" in text:
            account = "D.1.04"
        else:
            account = "D.1.03"
        confidence = 0.96
        reason = "Payroll/benefits expense column."
    elif "parish" in text and "guest" in text and "stipend" in text:
        account, confidence = "C.2.04", 0.98
        reason = "Combined parish and guest priest stipend column."
    elif "parish priest" in text and "stipend" in text:
        account, confidence = "C.2.01", 0.98
        reason = "Parish priest stipend column."
    elif "parochial vicar" in text and "stipend" in text:
        account, confidence = "C.2.02", 0.98
        reason = "Parochial vicar stipend column."
    elif "guest priest" in text and "stipend" in text:
        account, confidence = "C.2.03", 0.98
        reason = "Mass stipend expense column."
    elif "priest share" in text and "bishop" not in text:
        if "sacraments" in text:
            account = "C.1.01"
        elif "mass intentions" in text:
            account = "C.1.02"
        elif "confirmation" in text:
            account = "C.1.03"
        else:
            account = "C.1.05"
        confidence = 0.96
        reason = "Priest share expense column."
    elif "confirmation" in text and "minister-priest" in text:
        account, confidence = "C.1.04", 0.96
        reason = "Confirmation minister-priest expense column."
    elif "confirmation" in text and ("minister-others" in text or "minister others" in text):
        account, confidence = "C.3.01", 0.96
        reason = "Confirmation minister-others expense column."
    elif "other pastoral" in text:
        account, confidence = "C.3.02", 0.9
        reason = "Other pastoral expense column."
    elif "construction" not in section and any(k in text for k in ["food", "groceries", "meetings", "gasoline", "transportation", "office", "security", "liturgical", "real properties", "repairs", "charitable", "subscriptions", "hospital", "medicine", "other expenses"]):
        if "food" in text or "groceries" in text:
            account = "D.5.01"
        elif "meetings" in text:
            account = "D.5.02"
        elif "gasoline" in text:
            account = "D.5.03"
        elif "transportation" in text:
            account = "D.5.04"
        elif "office" in text:
            account = "D.5.05"
        elif "security" in text:
            account = "D.5.06"
        elif "liturgical" in text:
            account = "D.5.07"
        elif "real properties" in text:
            account = "D.5.08"
        elif "repairs" in text:
            account = "D.5.09"
        elif "charitable" in text:
            account = "D.5.10"
        elif "subscriptions" in text:
            account = "D.5.11"
        elif "hospital" in text or "medicine" in text:
            account = "D.5.12"
        elif "other expenses" in text:
            account = "D.5.13"
        confidence = 0.94
        reason = "Other parish/rectory expense column."
    elif "construction" in section and "donations" in text:
        account, confidence = "E.1.01", 0.96
        reason = "Construction receipt column."
    elif "construction" in section and "borrowings from banks" in text and "payment" not in text:
        account, confidence = "E.1.02", 0.96
        reason = "Construction receipt column."
    elif "construction" in section and "borrowings from other parishes" in text and "payment" not in text:
        account, confidence = "E.1.03", 0.96
        reason = "Construction receipt column."
    elif "construction" in section and "other receipts" in text:
        account, confidence = "E.1.04", 0.94
        reason = "Construction receipt column."
    elif "construction" in section and "labor" in text:
        account, confidence = "E.2.01", 0.96
        reason = "Construction expense column."
    elif "construction" in section and "payment of borrowings from banks" in text:
        account, confidence = "E.2.02", 0.96
        reason = "Construction expense column."
    elif "construction" in section and "payment of borrowings from other parishes" in text:
        account, confidence = "E.2.03", 0.96
        reason = "Construction expense column."
    elif "construction" in section and "other expenses" in text:
        account, confidence = "E.2.04", 0.94
        reason = "Construction expense column."
    elif "weekday collections" in text:
        account, confidence = ("B.1.05", 0.96) if "envelopes" in text else ("B.1.01", 0.98)
        reason = "Weekday mass collection."
    elif "sunday collections" in text and "bible" not in text:
        account, confidence = ("B.1.06", 0.96) if "envelopes" in text else ("B.1.02", 0.98)
        reason = "Sunday mass collection."
    elif "saturday anticipated" in text:
        account, confidence = "B.1.03", 0.98
        reason = "Saturday anticipated mass collection."
    elif any(k in text for k in ["rentals", "mortuary", "columbary", "kandilaan", "donation boxes", "envelopes", "other sources"]):
        if "rentals" in text:
            account = "B.2.01"
        elif "mortuary" in text or "columbary" in text:
            account = "B.2.02"
        elif "kandilaan" in text:
            account = "B.2.03"
        elif "donation boxes" in text:
            account = "B.2.04"
        elif "envelopes" in text:
            account = "B.2.05"
        else:
            account = "B.2.06"
        confidence = 0.95
        reason = "Other collections subject to collection grouping."
    elif "interest income" in text:
        account, confidence = "B.3.02", 0.98
        reason = "Interest income column."
    elif "subsidy" in text:
        account, confidence = "B.3.03", 0.98
        reason = "Subsidy inflow column."
    elif "second collections" in text:
        account, confidence, status = "B.3.05", 0.9, "needs_review"
        reason = "Second Collections appears in different source sections across years; review required."
    elif "special collections" in text and "remittances" not in section:
        account, confidence = "B.3.04", 0.94
        reason = "Special collections receipt column."
    elif "donations" in text and "construction" not in section:
        account, confidence = "B.3.01", 0.96
        reason = "Donation receipt column."
    elif "other receipts" in text and "construction" not in section:
        account, confidence = "B.3.07", 0.92
        reason = "Other receipt column."
    elif "remittances" in section and any(k in text for k in ["diocese share", "diocese fund", "pension fund", "progressive tax", "5% tax collections"]):
        if "sacraments" in text:
            account = "F.1.01"
        elif "diocese fund" in text:
            account = "F.1.02"
        elif "pension fund" in text:
            account = "F.1.03"
        elif "progressive tax" in text:
            account = "F.1.04"
        else:
            account = "F.1.05"
        confidence = 0.96
        reason = "Remittance to diocese column."
    elif "bishop" in text:
        if "minister-bishop" in text:
            account = "F.2.02"
        elif "others" in text:
            account = "F.2.03"
        else:
            account = "F.2.01"
        confidence = 0.94
        reason = "Bishop's fund share column."
    elif "remittances" in section and (
        any(
            k in text
            for k in [
                "pro-nigritis",
                "sancta infantia",
                "nat'l bible",
                "national bible",
                "special coll",
                "special collection",
                "2nd collection",
                "second collection",
                "national youth day",
                "simbang gabi",
                "catechetical",
            ]
        )
        or "other collections" in text
    ):
        account, confidence = "F.3.01", 0.9
        reason = "Special collection remittance column."
    elif text.strip() == "other collections":
        account, confidence = "F.3.01", 0.88
        status = "needs_review"
        reason = "Standalone Other Collections is treated as special collection remittance; review because the source header is sparse."
    elif "confirmation" in text and "bishop" not in text and "diocese" not in text and "pension" not in text and "minister" not in text:
        account, confidence = "A.2.01", 0.96
        reason = "Confirmation receipt column."
    elif "charge over" in text or "over/above" in text:
        account, confidence = "A.1.11", 0.78
        status = "needs_review"
        reason = "Charge over/above amount; review because source placement changes by year."
    elif any(k in text for k in ["baptism", "wedding", "funeral", "certificates", "marriage banns", "permits"]):
        account, confidence = sacrament_code, 0.72
        status = "needs_review"
        reason = "Sacrament column detected; review whether this is a final amount or a checking/detail value."
    elif "ending cash balance" in text or "construction fund net receipts" in text:
        account, field, confidence = None, "balance_or_reconciliation", 0.8
        status, rule = "needs_review", "memo"
        reason = "Balance/reconciliation field; review before treating as a line item."

    if account:
        if status != "needs_review":
            status = "approved" if confidence >= 0.86 else "needs_review"
        return account, account_names.get(account), field, confidence, status, reason, rule, combined
    return None, None, field, confidence, status, reason, rule, combined


def _seed_account_fallback() -> list[dict[str, Any]]:
    seed_path = Path(__file__).resolve().parents[4] / "supabase" / "portable_seeds" / "003_iafr_canonical_accounts.sql"
    if not seed_path.exists():
        return []
    rows: list[dict[str, Any]] = []
    for line in seed_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped.startswith("('"):
            continue
        values = [value.replace("''", "'") for value in re.findall(r"'((?:''|[^'])*)'", stripped)]
        if len(values) < 9:
            continue
        rows.append(
            {
                "section_code": values[0],
                "subsection_code": values[1],
                "account_code": values[2],
                "account_name": values[3],
                "account_type": values[4],
                "classification": values[5],
                "parent_account_code": values[6],
                "source_template": values[7],
                "source_sheet_name": values[8],
                "is_active": True,
            }
        )
    return rows


def _get_account_lookup() -> tuple[dict[str, str], dict[str, dict[str, Any]]]:
    fallback_rows = _seed_account_fallback()
    fallback_by_code = {row["account_code"]: row for row in fallback_rows}
    try:
        result = get_table("parishes", "iafr_account_titles").select("*").eq("is_active", True).execute()
        rows = result.data or []
        by_code = {**fallback_by_code, **{row["account_code"]: row for row in rows}}
        return {code: row["account_name"] for code, row in by_code.items()}, by_code
    except Exception:
        return {code: row["account_name"] for code, row in fallback_by_code.items()}, fallback_by_code


def list_parish_institutions() -> list[dict[str, Any]]:
    try:
        rows = (
            get_table("diocese", "institutions")
            .select("id, name, institution_code, vicariate, class, institution_type")
            .eq("institution_type", "parish")
            .is_("deleted_at", "null")
            .execute()
            .data
            or []
        )
    except Exception:
        return []
    return sorted(
        [
            {
                "id": row["id"],
                "name": row.get("name"),
                "institution_code": row.get("institution_code"),
                "vicariate": row.get("vicariate"),
                "class": row.get("class"),
            }
            for row in rows
            if row.get("id") and row.get("name")
        ],
        key=lambda row: (row.get("name") or "", row.get("institution_code") or ""),
    )


def _load_institution_matchers() -> tuple[dict[str, list[dict[str, Any]]], list[dict[str, Any]], dict[tuple[str | None, str], str]]:
    try:
        institutions = (
            get_table("diocese", "institutions")
            .select("id, name, institution_type, institution_code")
            .eq("institution_type", "parish")
            .is_("deleted_at", "null")
            .execute()
            .data
            or []
        )
    except Exception:
        institutions = []

    try:
        aliases = (
            get_table("operations", "parish_import_aliases")
            .select("source_code, normalized_source_name, institution_id, status")
            .eq("status", "approved")
            .is_("deleted_at", "null")
            .execute()
            .data
            or []
        )
    except Exception:
        # The alias table is created by the pusher staging migration. Preview
        # matching must still work before that migration is applied.
        aliases = []

    exact: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in institutions:
        if row.get("name"):
            exact[normalize_name(row.get("name"))].append(row)
    alias_lookup = {
        (row.get("source_code"), row.get("normalized_source_name")): row["institution_id"]
        for row in aliases
        if row.get("normalized_source_name")
    }
    return exact, institutions, alias_lookup


def _parish_review_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    review: dict[tuple[str | None, str], dict[str, Any]] = {}
    for row in rows:
        key = (row.get("parish_code"), normalize_name(row.get("parish_name")))
        if key not in review:
            review[key] = {
                "sourceCode": row.get("parish_code"),
                "sourceName": row.get("parish_name"),
                "institutionId": row.get("institution_id"),
                "matchStatus": row.get("institution_match_status"),
                "matchName": row.get("institution_match_name"),
                "confidence": row.get("institution_match_confidence"),
                "validationStatus": row.get("validation_status"),
                "monthCount": 0,
                "rowCount": 0,
                "_months": set(),
            }
        item = review[key]
        item["rowCount"] += 1
        item["_months"].add(row.get("reporting_month"))
        item["monthCount"] = len(item["_months"])
        if row.get("validation_status") == "blocked":
            item["validationStatus"] = "blocked"
        elif row.get("validation_status") == "warning" and item["validationStatus"] != "blocked":
            item["validationStatus"] = "warning"
    for item in review.values():
        item.pop("_months", None)
    return sorted(review.values(), key=lambda item: (item["validationStatus"] != "blocked", item["sourceName"] or ""))


def _match_institution(
    code: str | None,
    name: str,
    exact: dict[str, list[dict[str, Any]]],
    institutions: list[dict[str, Any]],
    aliases: dict[tuple[str | None, str], str],
) -> dict[str, Any]:
    normalized = normalize_name(name)
    source_variants = [normalized]
    if "," in name:
        source_variants.append(normalize_name(name.split(",", 1)[0]))
    alias_id = aliases.get((code, normalized)) or aliases.get((None, normalized))
    if alias_id:
        return {"institution_id": alias_id, "status": "matched", "confidence": 1.0, "match_name": name}
    for variant in source_variants:
        if variant in exact:
            candidates = exact[variant]
            status = "matched" if len(candidates) == 1 else "suggested"
            candidate = candidates[0]
            return {"institution_id": candidate["id"], "status": status, "confidence": 1.0, "match_name": candidate.get("name")}

    best = None
    best_score = 0.0
    second_score = 0.0
    for row in institutions:
        candidate = row.get("name") or ""
        candidate_norm = normalize_name(candidate)
        score = max(_name_similarity(variant, candidate_norm) for variant in source_variants)
        if score > best_score:
            second_score = best_score
            best = row
            best_score = score
        elif score > second_score:
            second_score = score
    if best and best_score >= 0.86:
        if best_score - second_score < 0.03:
            return {"institution_id": best["id"], "status": "suggested", "confidence": round(best_score, 2), "match_name": best.get("name")}
        return {"institution_id": best["id"], "status": "suggested", "confidence": round(best_score, 2), "match_name": best.get("name")}
    return {"institution_id": None, "status": "unmatched", "confidence": round(best_score, 2), "match_name": best.get("name") if best else None}


def _is_parish_data_row(code: Any, parish_name: Any) -> tuple[bool, str | None, str]:
    name = str(parish_name or "").strip()
    if not name:
        return False, None, ""

    normalized = normalize_name(name)
    if not normalized or normalized in {"sub total", "subtotal", "total"} or normalized.startswith("district "):
        return False, None, ""

    if isinstance(code, str):
        code_text = code.strip()
        if code_text.startswith("D") and "-" in code_text:
            return True, code_text, name

    if isinstance(code, (int, float)) and float(code).is_integer() and int(code) > 0:
        return True, None, name

    return False, None, ""


def _source_columns(ws, account_names: dict[str, str]) -> list[SourceColumn]:
    grid = _header_grid(ws)
    columns: list[SourceColumn] = []
    occurrences: dict[str, int] = defaultdict(int)
    for idx in range(3, ws.max_column + 1):
        parts = []
        for row in (3, 4, 5, 6):
            value = _clean_header(grid[row - 1][idx - 1])
            if value and value not in parts:
                parts.append(value)
        if not _is_financial_value_column(parts):
            continue
        column_letter = get_column_letter(idx)
        header = parts[-1] if parts else column_letter
        header_path = " > ".join(parts)
        base_key = _source_column_key(parts, 1)
        occurrences[base_key] += 1
        source_key = _source_column_key(parts, occurrences[base_key])
        account, account_name, field, confidence, status, reason, rule, combined = _suggest_mapping(parts, account_names)
        columns.append(
            SourceColumn(
                key=source_key,
                column=column_letter,
                index=idx,
                source_section=parts[0] if parts else "",
                source_header=header,
                header_path=header_path,
                suggested_account_code=account,
                suggested_account_name=account_name,
                suggested_field=field,
                confidence=round(confidence, 2),
                status=status,
                reason=reason,
                aggregation_rule=rule,
                is_combined=combined,
            )
        )
    return columns


def _parse_workbook(file_name: str, file_bytes: bytes) -> dict[str, Any]:
    account_names, _ = _get_account_lookup()
    exact, institutions, aliases = _load_institution_matchers()
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    month_sheets = []
    for sheet_name in wb.sheetnames:
        parsed = _month_year_from_sheet(sheet_name)
        if parsed:
            month_sheets.append((sheet_name, parsed[0], parsed[1]))
    if not month_sheets:
        raise ValueError("No monthly sheets found. Expected names like 'January 2024'.")

    sheet_columns: dict[str, list[SourceColumn]] = {}
    columns_by_key: dict[str, SourceColumn] = {}
    for sheet_name, _, _ in month_sheets:
        active_columns = _source_columns(wb[sheet_name], account_names)
        sheet_columns[sheet_name] = active_columns
        for col in active_columns:
            columns_by_key.setdefault(col.key, col)
    columns = list(columns_by_key.values())
    years = sorted({year for _, _, year in month_sheets})
    detected_year = years[0] if len(years) == 1 else None

    rows = []
    parish_keys: set[str] = set()
    zero_filled = 0
    issues_by_type = defaultdict(int)

    for sheet_name, month_num, year in month_sheets:
        ws = wb[sheet_name]
        active_columns = sheet_columns[sheet_name]
        for row_number in range(1, ws.max_row + 1):
            code = ws.cell(row_number, 1).value
            parish_name = ws.cell(row_number, 2).value
            is_parish_row, parish_code, parish_name_text = _is_parish_data_row(code, parish_name)
            if not is_parish_row:
                continue

            parish_keys.add(parish_code or normalize_name(parish_name_text))
            raw_values: dict[str, Any] = {}
            cleaned_values: dict[str, float] = {}
            row_issues: list[dict[str, Any]] = []
            row_zeroes = 0

            for col in active_columns:
                cell_value = ws.cell(row_number, col.index).value
                raw_values[col.key] = cell_value
                value, had_error = parse_money(cell_value)
                cleaned_values[col.key] = round(value, 2)
                if cell_value in (None, ""):
                    row_zeroes += 1
                if had_error:
                    row_issues.append(
                        {
                            "severity": "error",
                            "type": "invalid_numeric",
                            "message": f"{col.source_header} contains a non-numeric or formula-error value.",
                            "column": col.column,
                        }
                    )

            zero_filled += row_zeroes
            match = _match_institution(parish_code, parish_name_text, exact, institutions, aliases)
            if match["status"] == "unmatched":
                row_issues.append(
                    {
                        "severity": "blocker",
                        "type": "unmatched_parish",
                        "message": "Parish could not be matched to an institution.",
                    }
                )
            elif match["status"] == "suggested":
                row_issues.append(
                    {
                        "severity": "warning",
                        "type": "suggested_parish_match",
                        "message": f"Suggested match: {match['match_name']} ({match['confidence']}).",
                    }
                )

            validation_status = "blocked" if any(i["severity"] in ("blocker", "error") for i in row_issues) else "warning" if row_issues else "ready"
            for issue in row_issues:
                issues_by_type[issue["type"]] += 1

            rows.append(
                {
                    "sheet_name": sheet_name,
                    "source_row_number": row_number,
                    "parish_code": parish_code,
                    "parish_name": parish_name_text,
                    "institution_id": match["institution_id"],
                    "institution_match_status": match["status"],
                    "institution_match_name": match["match_name"],
                    "institution_match_confidence": match["confidence"],
                    "reporting_month": month_num,
                    "reporting_month_label": MONTH_TO_SHORT[month_num],
                    "reporting_year": year,
                    "validation_status": validation_status,
                    "raw_values": raw_values,
                    "cleaned_values": cleaned_values,
                    "issues": row_issues,
                    "zero_filled_fields": row_zeroes,
                }
            )

    columns_payload = [col.__dict__ for col in columns]
    ready = sum(1 for row in rows if row["validation_status"] == "ready")
    warnings = sum(1 for row in rows if row["validation_status"] == "warning")
    blocked = sum(1 for row in rows if row["validation_status"] == "blocked")
    mapping_review_count = sum(1 for col in columns if col.status == "needs_review")
    summary = {
        "fileName": file_name,
        "detectedYear": detected_year,
        "monthCount": len(month_sheets),
        "parishCount": len(parish_keys),
        "extractedRowCount": len(rows),
        "readyRowCount": ready,
        "warningRowCount": warnings,
        "blockedRowCount": blocked,
        "columnCount": len(columns),
        "mappingReviewCount": mapping_review_count,
        "zeroFilledCellCount": zero_filled,
        "issuesByType": dict(issues_by_type),
    }
    return {"summary": summary, "columns": columns_payload, "rows": rows, "parishMatches": _parish_review_rows(rows)}


def _persist_preview(file_name: str, file_hash: str, parsed: dict[str, Any], uploaded_by: str | None) -> str | None:
    try:
        summary = parsed["summary"]
        batch = (
            get_table("operations", "financial_push_batches")
            .insert(
                {
                    "source_file_name": file_name,
                    "source_file_hash": file_hash,
                    "detected_year": summary["detectedYear"],
                    "status": "validated",
                    "uploaded_by": uploaded_by,
                    "month_count": summary["monthCount"],
                    "parish_count": summary["parishCount"],
                    "extracted_row_count": summary["extractedRowCount"],
                    "valid_row_count": summary["readyRowCount"],
                    "warning_count": summary["warningRowCount"] + summary["mappingReviewCount"],
                    "error_count": summary["blockedRowCount"],
                    "zero_filled_cell_count": summary["zeroFilledCellCount"],
                    "summary": summary,
                }
            )
            .execute()
        )
        batch_id = batch.data[0]["id"]

        for col in parsed["columns"]:
            try:
                get_table("operations", "financial_import_column_map").upsert(
                    {
                        "template_year": summary["detectedYear"],
                        "source_header": col["source_header"],
                        "source_section": col["source_section"],
                        "source_column": col["column"],
                        "canonical_account_code": col["suggested_account_code"],
                        "canonical_field": col["suggested_field"],
                        "aggregation_rule": col["aggregation_rule"],
                        "is_combined": col["is_combined"],
                        "confidence": col["confidence"],
                        "status": "approved" if col["status"] == "approved" else "needs_review",
                        "notes": col["reason"],
                    },
                    on_conflict="template_year,source_header,source_section,source_column",
                ).execute()
            except Exception:
                pass

        payload = [
            {
                "push_batch_id": batch_id,
                "sheet_name": row["sheet_name"],
                "source_row_number": row["source_row_number"],
                "parish_code": row["parish_code"],
                "parish_name": row["parish_name"],
                "institution_id": row["institution_id"],
                "reporting_month": row["reporting_month"],
                "reporting_year": row["reporting_year"],
                "validation_status": row["validation_status"],
                "raw_values": row["raw_values"],
                "cleaned_values": row["cleaned_values"],
                "mapped_values": {},
                "issues": row["issues"],
                "zero_filled_fields": row["zero_filled_fields"],
            }
            for row in parsed["rows"]
        ]
        for i in range(0, len(payload), 500):
            get_table("operations", "financial_push_rows").insert(payload[i : i + 500]).execute()
        return batch_id
    except Exception:
        return None


def validate_files(files: list[dict[str, str]], uploaded_by: str | None = None) -> dict[str, Any]:
    results = []
    for file in files:
        file_name = file["name"]
        raw = base64.b64decode(file["contentBase64"])
        file_hash = hashlib.sha256(raw).hexdigest()
        parsed = _parse_workbook(file_name, raw)
        batch_id = _persist_preview(file_name, file_hash, parsed, uploaded_by)
        results.append(
            {
                "batchId": batch_id,
                "fileName": file_name,
                "summary": parsed["summary"],
                "columns": parsed["columns"],
                "parishMatches": parsed["parishMatches"],
                "sampleRows": parsed["rows"][:12],
                "persistenceStatus": "persisted" if batch_id else "preview_only",
            }
        )
    return {"files": results}


def list_accounts() -> list[dict[str, Any]]:
    _, rows = _get_account_lookup()
    if rows:
        return sorted(rows.values(), key=lambda r: r.get("account_code") or "")
    names, _ = _get_account_lookup()
    return [{"account_code": code, "account_name": name, "section_code": code[0]} for code, name in sorted(names.items())]


def create_canonical_account(payload: dict[str, Any]) -> dict[str, Any]:
    account = {
        "section_code": payload["sectionCode"],
        "subsection_code": payload.get("subsectionCode") or None,
        "account_code": payload["accountCode"],
        "account_name": payload["accountName"],
        "account_type": payload["accountType"],
        "classification": payload.get("classification") or None,
        "source_template": f"PUSHER_{payload.get('effectiveYear') or 'HISTORICAL'}",
        "source_sheet_name": payload.get("sourceHeader") or None,
        "is_active": True,
    }
    result = get_table("parishes", "iafr_account_titles").insert(account).execute()
    created = result.data[0]
    try:
        get_table("operations", "canonical_account_requests").insert(
            {
                "source_header": payload.get("sourceHeader") or payload["accountName"],
                "source_section": payload.get("sourceSection"),
                "suggested_account_code": payload["accountCode"],
                "suggested_account_name": payload["accountName"],
                "suggested_section_code": payload["sectionCode"],
                "suggested_subsection_code": payload.get("subsectionCode"),
                "suggested_account_type": payload["accountType"],
                "suggested_classification": payload.get("classification"),
                "effective_year": payload.get("effectiveYear"),
                "reason": payload.get("reason"),
                "status": "created",
                "requested_by": payload.get("requestedBy"),
                "reviewed_by": payload.get("requestedBy"),
                "reviewed_at": datetime.now(timezone.utc).isoformat(),
                "created_account_title_id": created["id"],
            }
        ).execute()
    except Exception:
        pass
    return created


def _account_payload(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        "section_code": payload["sectionCode"],
        "subsection_code": payload.get("subsectionCode") or None,
        "account_code": payload["accountCode"],
        "account_name": payload["accountName"],
        "account_type": payload["accountType"],
        "classification": payload.get("classification") or None,
        "source_template": f"PUSHER_{payload.get('effectiveYear') or 'HISTORICAL'}",
        "source_sheet_name": payload.get("sourceHeader") or None,
        "is_active": True,
    }


def update_canonical_account(payload: dict[str, Any]) -> dict[str, Any]:
    account_id = payload.get("id")
    if not account_id:
        raise ValueError("Canonical account id is required.")
    result = (
        get_table("parishes", "iafr_account_titles")
        .update(_account_payload(payload))
        .eq("id", account_id)
        .execute()
    )
    if not result.data:
        raise ValueError("Canonical account was not found.")
    return result.data[0]


def deactivate_canonical_account(account_id: str) -> dict[str, Any]:
    if not account_id:
        raise ValueError("Canonical account id is required.")
    result = (
        get_table("parishes", "iafr_account_titles")
        .update({"is_active": False, "deleted_at": datetime.now(timezone.utc).isoformat()})
        .eq("id", account_id)
        .execute()
    )
    if not result.data:
        raise ValueError("Canonical account was not found.")
    return result.data[0]


def _get_existing_record(institution_id: str, year: int, month_label: str) -> str | None:
    result = (
        get_table("parishes", "financial_records")
        .select("id")
        .eq("institution_id", institution_id)
        .eq("year", year)
        .eq("month", month_label)
        .eq("is_current_version", True)
        .is_("deleted_at", "null")
        .maybe_single()
        .execute()
    )
    data = getattr(result, "data", None)
    return data["id"] if data else None


def _ensure_record(institution_id: str, batch_id: str, year: int, month: int, import_mode: str) -> tuple[str | None, str]:
    month_label = MONTH_TO_SHORT[month]
    existing_id = _get_existing_record(institution_id, year, month_label)
    now = datetime.now(timezone.utc).isoformat()
    if existing_id and import_mode == "skip_existing":
        return None, "skipped_existing"
    if existing_id:
        get_table("parishes", "financial_records").update(
            {
                "submission_batch_id": None,
                "status": "submitted",
                "submitted_at": now,
                "record_timestamp": now,
                "beginning_balance": 0,
            }
        ).eq("id", existing_id).execute()
        return existing_id, "updated"
    inserted = (
        get_table("parishes", "financial_records")
        .insert(
            {
                "institution_id": institution_id,
                "month": month_label,
                "year": year,
                "status": "submitted",
                "submitted_at": now,
                "record_timestamp": now,
                "beginning_balance": 0,
            }
        )
        .execute()
    )
    inserted_data = getattr(inserted, "data", None) or []
    if not inserted_data:
        raise ValueError("Could not create financial record for the parish-month being committed.")
    return inserted_data[0]["id"], "inserted"


def start_commit_batch(batch_id: str, import_mode: str, committed_by: str | None = None) -> dict[str, Any]:
    batch = (
        get_table("operations", "financial_push_batches")
        .select("id, status, upload_mode")
        .eq("id", batch_id)
        .maybe_single()
        .execute()
    )
    data = getattr(batch, "data", None)
    if not data:
        raise ValueError("PUSHER batch was not found.")
    if data.get("status") == "committed":
        return {"batchId": batch_id, "status": "committed", "alreadyFinished": True}
    if data.get("status") == "committing":
        return {"batchId": batch_id, "status": "already_committing", "alreadyRunning": True}
    if _count_push_rows(batch_id) == 0:
        raise ValueError("No parish rows were staged for this file. Validate the workbook again before pushing.")
    get_table("operations", "financial_push_batches").update(
        {"status": "committing", "import_mode": import_mode, "committed_by": committed_by}
    ).eq("id", batch_id).execute()
    return {"batchId": batch_id, "status": "committing", "alreadyRunning": False}


def mark_commit_failed(batch_id: str, message: str) -> None:
    try:
        get_table("operations", "financial_push_batches").update(
            {
                "status": "failed",
                "summary": {
                    "commitError": message,
                    "failedAt": datetime.now(timezone.utc).isoformat(),
                },
            }
        ).eq("id", batch_id).execute()
    except Exception:
        pass


def _count_push_rows(batch_id: str, status: str | None = None) -> int:
    query = (
        get_table("operations", "financial_push_rows")
        .select("id", count="exact")
        .eq("push_batch_id", batch_id)
        .is_("deleted_at", "null")
        .limit(1)
    )
    if status:
        query = query.eq("validation_status", status)
    result = query.execute()
    return int(getattr(result, "count", 0) or 0)


def _load_push_rows_for_commit(batch_id: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    page_size = 1000
    start = 0
    columns = (
        "id, parish_code, parish_name, institution_id, reporting_month, reporting_year, "
        "validation_status, source_row_number, cleaned_values, issues"
    )
    while True:
        page = (
            get_table("operations", "financial_push_rows")
            .select(columns)
            .eq("push_batch_id", batch_id)
            .is_("deleted_at", "null")
            .range(start, start + page_size - 1)
            .execute()
            .data
            or []
        )
        rows.extend(page)
        if len(page) < page_size:
            break
        start += page_size
    return rows


def get_commit_progress(batch_id: str) -> dict[str, Any]:
    batch = (
        get_table("operations", "financial_push_batches")
        .select("id, status, upload_mode, summary, committed_at, updated_at")
        .eq("id", batch_id)
        .maybe_single()
        .execute()
    )
    batch_data = getattr(batch, "data", None)
    if not batch_data:
        raise ValueError("PUSHER batch was not found.")

    total = _count_push_rows(batch_id)
    committed = _count_push_rows(batch_id, "committed")
    skipped = _count_push_rows(batch_id, "skipped")
    blocked = _count_push_rows(batch_id, "blocked")
    ready = _count_push_rows(batch_id, "ready")
    warning = _count_push_rows(batch_id, "warning")
    processed = committed + skipped + blocked
    percent = round((processed / total) * 100, 1) if total else 0
    summary = batch_data.get("summary") or {}

    return {
        "batchId": batch_id,
        "status": batch_data.get("status"),
        "uploadMode": batch_data.get("upload_mode"),
        "totalRows": total,
        "processedRows": processed,
        "committedRows": committed,
        "skippedRows": skipped,
        "blockedRows": blocked,
        "readyRows": ready,
        "warningRows": warning,
        "lineItemsCreated": summary.get("lineItemsCreated"),
        "percent": percent,
        "committedAt": batch_data.get("committed_at"),
        "updatedAt": batch_data.get("updated_at"),
        "error": summary.get("commitError"),
    }


def list_batch_rows(batch_id: str, status: str | None = None, limit: int = 500) -> dict[str, Any]:
    safe_limit = max(1, min(limit, 2000))
    query = (
        get_table("operations", "financial_push_rows")
        .select(
            "id, parish_code, parish_name, reporting_month, reporting_year, "
            "validation_status, financial_record_id, issues"
        )
        .eq("push_batch_id", batch_id)
        .is_("deleted_at", "null")
        .order("reporting_month")
        .order("parish_name")
        .limit(safe_limit)
    )
    if status:
        query = query.eq("validation_status", status)
    rows = query.execute().data or []
    total = _count_push_rows(batch_id, status)
    return {
        "batchId": batch_id,
        "status": status,
        "total": total,
        "rows": [
            {
                "id": row.get("id"),
                "parishCode": row.get("parish_code"),
                "parishName": row.get("parish_name"),
                "reportingMonth": row.get("reporting_month"),
                "reportingMonthLabel": MONTH_TO_SHORT.get(row.get("reporting_month")),
                "reportingYear": row.get("reporting_year"),
                "validationStatus": row.get("validation_status"),
                "financialRecordId": row.get("financial_record_id"),
                "issues": row.get("issues") or [],
                "reason": (
                    "Existing parish-month record was kept because Skip existing was selected."
                    if row.get("validation_status") == "skipped"
                    else "Row was not committed. Review parish match or validation issues."
                ),
            }
            for row in rows
        ],
    }


def commit_batch(
    batch_id: str,
    mappings: list[dict[str, Any]],
    parish_mappings: list[dict[str, Any]] | None = None,
    import_mode: str = "skip_existing",
    committed_by: str | None = None,
) -> dict[str, Any]:
    accounts = {row["account_code"]: row for row in list_accounts() if row.get("account_code")}
    mapping_by_key = {m["key"]: m for m in mappings}
    approved = {
        key: value
        for key, value in mapping_by_key.items()
        if value.get("action") == "map" and value.get("canonicalAccountCode") in accounts
    }
    if not approved:
        raise ValueError("No approved mappings were provided.")
    reviewed_parishes = {
        (mapping.get("sourceCode"), normalize_name(mapping.get("sourceName"))): mapping.get("institutionId")
        for mapping in parish_mappings or []
        if mapping.get("institutionId") and mapping.get("sourceName")
    }

    rows = _load_push_rows_for_commit(batch_id)
    if not rows:
        raise ValueError("No parish rows were staged for this file. Validate the workbook again before pushing.")

    committed = skipped = blocked = line_items_created = 0
    for row in rows:
        row_key = (row.get("parish_code"), normalize_name(row.get("parish_name")))
        reviewed_institution_id = reviewed_parishes.get(row_key)
        if reviewed_institution_id:
            row["institution_id"] = reviewed_institution_id
            if row["validation_status"] in ("blocked", "warning"):
                blocking_issues = [
                    issue
                    for issue in row.get("issues", [])
                    if issue.get("severity") in ("blocker", "error") and issue.get("type") != "unmatched_parish"
                ]
                row["validation_status"] = "blocked" if blocking_issues else "ready"
                get_table("operations", "financial_push_rows").update(
                    {
                        "institution_id": reviewed_institution_id,
                        "validation_status": row["validation_status"],
                    }
                ).eq("id", row["id"]).execute()
        if row["validation_status"] != "ready" or not row.get("institution_id"):
            blocked += 1
            get_table("operations", "financial_push_rows").update({"validation_status": "blocked"}).eq("id", row["id"]).execute()
            continue
        record_id, action = _ensure_record(row["institution_id"], batch_id, row["reporting_year"], row["reporting_month"], import_mode)
        if not record_id:
            skipped += 1
            get_table("operations", "financial_push_rows").update({"validation_status": "skipped"}).eq("id", row["id"]).execute()
            continue

        get_table("parishes", "iafr_line_items").delete().eq("financial_record_id", record_id).execute()
        items = []
        cleaned_values = row.get("cleaned_values") or {}
        for key, mapping in approved.items():
            account_code = mapping["canonicalAccountCode"]
            account = accounts[account_code]
            amount = float(cleaned_values.get(key) or 0)
            if round(amount, 2) == 0:
                continue
            source_header = mapping.get("sourceHeader") or key.split(":", 1)[-1]
            items.append(
                {
                    "financial_record_id": record_id,
                    "account_title_id": account.get("id"),
                    "section_code": account.get("section_code"),
                    "subsection_code": account.get("subsection_code"),
                    "item_label": source_header[:200],
                    "item_type": account.get("account_type"),
                    "amount": round(amount, 2),
                    "source_row_number": row.get("source_row_number"),
                    "source_label": source_header[:200],
                }
            )
        if items:
            for i in range(0, len(items), 500):
                get_table("parishes", "iafr_line_items").insert(items[i : i + 500]).execute()
            line_items_created += len(items)
        committed += 1
        get_table("operations", "financial_push_rows").update(
            {"validation_status": "committed", "financial_record_id": record_id, "mapped_values": approved}
        ).eq("id", row["id"]).execute()

    now = datetime.now(timezone.utc).isoformat()
    get_table("operations", "financial_push_batches").update(
        {
            "status": "committed",
            "upload_mode": "committed",
            "committed_at": now,
            "committed_by": committed_by,
            "summary": {
                "committedRows": committed,
                "skippedRows": skipped,
                "blockedRows": blocked,
                "lineItemsCreated": line_items_created,
            },
        }
    ).eq("id", batch_id).execute()

    return {
        "batchId": batch_id,
        "committedRows": committed,
        "skippedRows": skipped,
        "blockedRows": blocked,
        "lineItemsCreated": line_items_created,
    }
